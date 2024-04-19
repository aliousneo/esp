from __future__ import annotations
import json
from typing import ClassVar

from elasticsearch import AsyncElasticsearch
from elasticsearch.helpers import async_streaming_bulk
import pandas as pd
from sentence_transformers import SentenceTransformer


class AsyncElasticService:

    client: ClassVar[AsyncElasticsearch|None] = None
    products_info_json_filename: ClassVar[str] = 'all_data.json'
    products_info_csv_filename: ClassVar[str] = 'all_data.csv'
    inbound_data_excel_files: ClassVar[dict[str, str]] = {
        'products_descriptions': 'ProductList 2024.xlsx',
        'ai_search_patterns': 'Customer Data for AI Searching.xlsx'
    }
    encoding_model: ClassVar[SentenceTransformer] = SentenceTransformer('all-mpnet-base-v2')
    es_index_name: ClassVar[str] = 'vect_search_products'
    products_es_model: ClassVar[dict[str, dict[str, dict[str, str|int|bool]]]] = {
        'properties': {
            'product_key': {
                'type': 'text'
            },
            'description': {
                'type': 'text'
            },
            'search_patterns': {
                'type': 'text'
            },
            'search_vector': {
                'type': 'dense_vector',
                'dims': 768,
                'index': True,
                'similarity': 'l2_norm'
            }
        }
    }

    @classmethod
    async def connect(cls: type[AsyncElasticService]) -> None:
        cls.client = AsyncElasticsearch(
            cloud_id="somecloud",
            basic_auth=("some", "any")
        )

    @classmethod
    async def close_connection(cls: type[AsyncElasticService]) -> None:
        await cls.client.close()

    @classmethod
    async def search(cls, query: str) -> dict[str, list[dict[str, str]]]:
        query_input = cls.encoding_model.encode(query)
        es_query_dict = {
            'field': 'search_vector',
            'query_vector': query_input,
            # 'k': 2,
            # 'num_candidates': 500
        }

        result = await AsyncElasticService.client.search(
            index=cls.es_index_name,
            knn=es_query_dict,
            source=['product_key', 'description'],
            min_score=0.50
        )

        return {
            'search_results':[
                {
                    'product_key': item['_source']['product_key'],
                    'description': item['_source']['description']
                }
                for item in result['hits']['hits']
            ]
        }

    @staticmethod
    async def simply_iterate(products_list):
        for product in products_list:
            yield product

    @classmethod
    async def ingest_data_into_es_index(cls) -> None:
        '''Ingest data into ElasticSearch'''
        await cls.excel_to_json()
        await cls.json_to_csv()

        if not (await cls.client.indices.exists(index=cls.es_index_name)):
            await cls.client.indices.create(index=cls.es_index_name, mappings=cls.products_es_model)

        df = pd.read_csv(cls.products_info_csv_filename)

        # Add new field with vector for each product (combines all 'ai_search_patterns' encoded together)
        df['search_vector'] = df['search_patterns'].apply(lambda x: cls.encoding_model.encode(x))

        products_list = df.to_dict(orient='records')

        # Upload all products info to ElasticSearch index
        async for _ in async_streaming_bulk(
            client=cls.client,
            index=cls.es_index_name,
            actions=cls.simply_iterate(products_list)
        ):
            pass

    @classmethod
    async def json_to_csv(cls) -> None:
        # Load data into memory from json file
        with open(cls.products_info_json_filename, 'r', encoding='utf-8') as f:
            data = json.loads(f.read())
        # Normalize data
        df = pd.json_normalize(data)
        # Store data in csv file
        df.to_csv(cls.products_info_csv_filename, index=False, encoding='utf-8')

    @classmethod
    async def excel_to_json(cls) -> None:
        from openpyxl import load_workbook

        wb1 = load_workbook(filename=cls.inbound_data_excel_files['products_descriptions'])
        sheet_1 = wb1['Sheet1']

        all_products: dict[str, dict[str, str|list[str]]] = {}

        # Attention: 'max_col' and 'max_row' are the max numbers of col/raws with data on the sheet
        for row in sheet_1.iter_rows(min_row=2, max_col=2, max_row=14525):
            product_key = str(row[0].value)
            description = row[1].value
            all_products[product_key] = {
                'product_key': product_key,
                'description': description,
                'search_patterns': [description]
            }

        wb2 = load_workbook(filename=cls.inbound_data_excel_files['ai_search_patterns'])
        sheet_2 = wb2['Sheet1']

        # Attention: 'max_col' and 'max_row' are the max numbers of col/raws with data on the shee
        for row in sheet_2.iter_rows(min_row=2, max_col=2, max_row=5763):
            product_key: str = str(row[1].value)
            ai_search_pattern: str = str(row[0].value)
            iterable_sequence: list[str] = [ai_search_pattern, product_key]

            # case when there are multiple keys in the 'product_key' cell - start -
            if '+' in product_key:
                product_keys = row[1].value.split(' + ')
                for pk in product_keys:
                    try:
                        if all_products[pk].get('search_patterns') is not None:
                            if ai_search_pattern not in all_products[pk]['search_patterns']:
                                all_products[pk]['search_patterns'].append(ai_search_pattern)
                        else:
                            all_products[pk]['search_patterns'] = [ai_search_pattern]
                    except KeyError:
                        for product in all_products:
                            if product.startswith(product_key):
                                if all_products[product].get('search_patterns') is not None:
                                    for value in iterable_sequence:
                                        if value not in all_products[pk]['search_patterns']:
                                            all_products[product]['search_patterns'].append(value)
                                else:
                                    all_products[product]['search_patterns'] = iterable_sequence
                continue
            # case when there are multiple keys in the cell - end -

            try:
                if all_products[product_key].get('search_patterns') is not None:
                    if ai_search_pattern not in all_products[product_key]['search_patterns']:
                        all_products[product_key]['search_patterns'].append(ai_search_pattern)
                else:
                    all_products[product_key]['search_patterns'] = [ai_search_pattern]
            except KeyError:
                for product in all_products:
                    if product.startswith(product_key):
                        if all_products[product].get('search_patterns') is not None:
                            for item in iterable_sequence:
                                if item not in all_products[product]['search_patterns']:
                                    all_products[product]['search_patterns'].append(str(item))

                        else:
                            all_products[product]['search_patterns'] = iterable_sequence

        # Create list of all products info
        all_products_info_list = [v for _, v in all_products.items()]

        # Store products info in json file
        with open(cls.products_info_json_filename, 'w') as f:
            json.dump(all_products_info_list, f)
